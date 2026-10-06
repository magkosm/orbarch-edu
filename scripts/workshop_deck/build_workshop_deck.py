#!/usr/bin/env python3
"""Build the workshop deck "The Architect's Mind in Orbit" on the KTH PowerPoint template.

Usage:  python3 scripts/workshop_deck/build_workshop_deck.py <en|sv|el> <out.pptx>
Needs:  python-pptx and pymupdf, plus a WORKDIR with the private inputs (see README.md here):
          $WORKDIR/kth-template.pptx      the official KTH PowerPoint template (its own slides are dropped)
          $WORKDIR/assets/…               photos and logos (ESA/M. Wandt, NASA/R. Guidice, ISAE-SUPAERO, ESERO, KTH)
          $WORKDIR/kth/ppt/media/…, $WORKDIR/school/ppt/media/…  media unpacked from earlier KTH decks
        WORKDIR defaults to this script's directory. Texts: strings.json next to it; screenshots from
        public/educators/screenshots/<lang>/; QR codes from docs/workshop/materials/qr/.
"""
import os, sys, json, copy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
WORK = Path(os.environ.get('WORKDIR', HERE))
LANG = sys.argv[1] if len(sys.argv) > 1 else 'en'
OUT = Path(sys.argv[2]) if len(sys.argv) > 2 else HERE / f'slides-{LANG}.pptx'
STR = json.loads((HERE / 'strings.json').read_text(encoding='utf-8'))[LANG]

A = WORK / 'assets'
M = WORK / 'kth/ppt/media'
SCH = WORK / 'school/ppt/media'
SHOT = REPO / 'public/educators/screenshots' / LANG
QR = REPO / 'docs/workshop/materials/qr'

IMG = dict(
    destiny=A / 'image32.jpeg', destiny_small=A / 'image34.jpg', guidice=A / 'image37.jpg', mars=A / 'image38.jpg',
    wandt=A / 'image43.jpeg', iss=A / 'image44.jpg', hatch=A / 'image45.jpg', cupola=A / 'image46.jpg', casa=A / 'image47.jpeg',
    mdrs_crew=M / 'image41.png', mdrs_desert=SCH / 'image28.jpeg', nback=M / 'image52.png', pvt=SCH / 'image48.jpeg',
    matb_old=SCH / 'image49.jpeg', muninn=M / 'image39.png',
    kth=M / 'image16.png', esa=M / 'image70.png', isae=M / 'image71.png', snsa=M / 'image69.png', esero=A / 'esero-white.png',
)

KTH_BLUE, NAVY, LIGHT, SAND, SKY, WHITE, GREY = (RGBColor(0x00, 0x47, 0x91), RGBColor(0x00, 0x00, 0x61), RGBColor(0xDE, 0xF0, 0xFF),
                                                  RGBColor(0xEB, 0xE5, 0xE0), RGBColor(0x61, 0x98, 0xD2), RGBColor(0xFF, 0xFF, 0xFF),
                                                  RGBColor(0x55, 0x5F, 0x6B))
GREEN, RED, YELLOW = RGBColor(0x2E, 0x8B, 0x57), RGBColor(0xC0, 0x39, 0x2B), RGBColor(0xFF, 0xE8, 0x7A)

prs = Presentation(WORK / 'kth-template.pptx')
for sldId in list(prs.slides._sldIdLst):
    prs.part.drop_rel(sldId.rId)
    prs.slides._sldIdLst.remove(sldId)
L = {l.name.strip(): l for l in prs.slide_layouts}
SW, SH = prs.slide_width, prs.slide_height


# ---------------------------------------------------------------- helpers
def layout(name):
    return L[name]


def ph(slide, idx):
    for p in slide.placeholders:
        if p.placeholder_format.idx == idx:
            return p
    raise KeyError(idx)


def kill_ph(slide, idx):
    for p in list(slide.placeholders):
        if p.placeholder_format.idx == idx:
            p._element.getparent().remove(p._element)


def set_text(tf, items, size=None, color=None, bold_first=False, align=None, space_after=6):
    """items: str or list of (text, level) / str; '**x**' marks bold runs."""
    if isinstance(items, str):
        items = [items]
    tf.clear()
    first = True
    for it in items:
        text, level = (it, 0) if isinstance(it, str) else it
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        if align: p.alignment = align
        p.space_after = Pt(space_after)
        parts = text.split('**')
        for i, part in enumerate(parts):
            if not part: continue
            r = p.add_run(); r.text = part
            if i % 2 == 1: r.font.bold = True
            if size: r.font.size = Pt(size)
            if color: r.font.color.rgb = color
    return tf


def textbox(slide, x, y, w, h, items, size=16, color=None, bold=False, align=None, anchor=None, font=None, space_after=4, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.03)
    if anchor: tf.vertical_anchor = anchor
    set_text(tf, items, size=size, color=color, align=align, space_after=space_after)
    if bold:
        for p in tf.paragraphs:
            for r in p.runs: r.font.bold = True
    if font:
        for p in tf.paragraphs:
            for r in p.runs: r.font.name = font
    return tb


def rect(slide, x, y, w, h, fill, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, transparency=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(1)
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    if transparency is not None:
        from pptx.oxml.ns import qn
        sf = s.fill._xPr.find(qn('a:solidFill'))
        clr = sf[0]
        a = clr.makeelement(qn('a:alpha'), {'val': str(int((100 - transparency) * 1000))}); clr.append(a)
    return s


def shape_text(s, items, size=16, color=WHITE, bold=False, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE):
    tf = s.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.1)
    set_text(tf, items, size=size, color=color, align=align, space_after=2)
    if bold:
        for p in tf.paragraphs:
            for r in p.runs: r.font.bold = True


def picture(slide, path, x, y, w=None, h=None):
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)


def fit_picture(slide, path, x, y, w, h):
    """Place picture inside a box keeping aspect ratio, centred."""
    import pymupdf
    pix = pymupdf.Pixmap(str(path)); ar = pix.width / pix.height
    if w / h > ar: ww, hh = h * ar, h
    else: ww, hh = w, w / ar
    return picture(slide, path, x + (w - ww) / 2, y + (h - hh) / 2, ww, hh)


def credit(slide, text, x=0.66, y=6.95, w=8, color=GREY, size=9):
    textbox(slide, x, y, w, 0.3, text, size=size, color=color)


def notes(slide, text):
    if text: slide.notes_slide.notes_text_frame.text = text


def title(slide, text, idx=0):
    p = ph(slide, idx); p.text_frame.text = text; return p


def arrow(slide, x1, y1, x2, y2, color=KTH_BLUE, width=2.5):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(width)
    from pptx.oxml.ns import qn
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'}); ln.append(tail)
    return c


def qr_slide(S, qr_key, shot, bullets, note=None, bullet_size=18):
    """Headline-only layout: bullets left, screenshot centre, QR right."""
    s = prs.slides.add_slide(layout('Headline only'))
    title(s, S['title'])
    textbox(s, 0.66, 2.1, 4.3, 4.6, [(b, 0) for b in bullets], size=bullet_size, color=NAVY, space_after=10)
    if shot:
        pic = fit_picture(s, SHOT / f'{shot}.png', 5.2, 2.1, 4.6, 3.3)
        pic.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDA); pic.line.width = Pt(0.75)
    qx = 10.05 if shot else 7.6
    r = rect(s, qx, 2.1, 2.62, 3.35, WHITE, line=RGBColor(0xC8, 0xD0, 0xDA), shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    picture(s, QR / f'{qr_key}-{LANG}.png', qx + 0.16, 2.2, 2.3, 2.3)
    textbox(s, qx, 4.55, 2.62, 0.8, [STR['scan'], f"**/{qr_key if qr_key != 'matb-2min' else '2min'}**"], size=11, color=NAVY, align=PP_ALIGN.CENTER)
    url = f"magkosm.github.io/orbarch-edu{'/' + ('2min' if qr_key == 'matb-2min' else qr_key) if qr_key != 'hub' else ''}?lng={LANG}"
    textbox(s, 5.2, 5.55, 7.5, 0.4, url, size=11, color=GREY)
    if shot:
        textbox(s, 5.2, 5.9, 7.5, 0.8, S.get('caption', ''), size=12, color=GREY)
    notes(s, note or S.get('notes', ''))
    return s


# ---------------------------------------------------------------- slides
T = STR['slides']

# 1 Title — full-bleed Destiny-lab photo
s = prs.slides.add_slide(layout('Empty'))
picture(s, IMG['destiny_small'], 0, 0, w=13.333)
rect(s, 0, 0, 13.333, 7.5, NAVY, transparency=45)
textbox(s, 0.8, 2.2, 11.7, 1.8, T['1']['title'], size=48, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
textbox(s, 0.8, 4.05, 11.7, 0.7, T['1']['sub'], size=20, color=WHITE)
textbox(s, 0.8, 4.75, 11.7, 0.5, T['1']['meta'], size=14, color=RGBColor(0xDE, 0xF0, 0xFF))
picture(s, IMG['esero'], 0.8, 6.3, h=0.7)
r = rect(s, 2.85, 6.22, 1.0, 0.86, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
picture(s, IMG['kth'], 2.95, 6.26, h=0.78)
credit(s, '© ESA – M. Wandt', x=9.9, y=7.05, w=3.3, color=RGBColor(0xDE, 0xF0, 0xFF), size=9)
notes(s, T['1']['notes'])

# 2 What is a space station — text and image
s = prs.slides.add_slide(layout('Text and image'))
title(s, T['2']['title'])
set_text(ph(s, 1).text_frame, [(b, 0) for b in T['2']['bullets']], size=18, space_after=10)
ph(s, 13).insert_picture(str(IMG['iss']))
credit(s, '© NASA / SpaceX', x=6.8, y=7.05, w=6, color=WHITE)
notes(s, T['2']['notes'])

# 3 History — timeline built from shapes
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['3']['title'])
stations = T['3']['stations']  # list of (name, year, where)
x0, x1, y = 1.2, 12.1, 4.2
arrow(s, x0 - 0.3, y, x1 + 0.4, y, color=KTH_BLUE, width=3)
n = len(stations)
for i, (name, year, where) in enumerate(stations):
    cx = x0 + i * (x1 - x0) / (n - 1)
    c = rect(s, cx - 0.22, y - 0.22, 0.44, 0.44, KTH_BLUE if i % 2 == 0 else NAVY, shape=MSO_SHAPE.OVAL)
    up = i % 2 == 0
    ty = y - 1.9 if up else y + 0.45
    textbox(s, cx - 1.15, ty, 2.3, 0.55, f"**{name}**", size=20, color=NAVY, align=PP_ALIGN.CENTER)
    textbox(s, cx - 1.15, ty + 0.5, 2.3, 0.45, year, size=16, color=KTH_BLUE, align=PP_ALIGN.CENTER)
    textbox(s, cx - 1.15, ty + 0.9, 2.3, 0.5, where, size=13, color=GREY, align=PP_ALIGN.CENTER)
textbox(s, 0.66, 6.35, 12, 0.5, T['3']['sub'], size=14, color=GREY)
notes(s, T['3']['notes'])

# 4 Machine-centric design
s = prs.slides.add_slide(layout('Text and image'))
title(s, T['4']['title'])
set_text(ph(s, 1).text_frame, [(b, 0) for b in T['4']['bullets']], size=18, space_after=10)
ph(s, 13).insert_picture(str(IMG['hatch']))
credit(s, '© ESA – M. Wandt', x=6.8, y=7.05, w=6, color=WHITE)
notes(s, T['4']['notes'])

# 5 Connection slide — 3 images with captions
s = prs.slides.add_slide(layout('Headline and 3 images, text'))
title(s, T['5']['title'])
imgs = [IMG['iss'], IMG['wandt'], IMG['mdrs_desert']] if LANG != 'el' else [IMG['iss'], IMG['wandt'], IMG['muninn']]
for idx, img in zip((13, 14, 15), imgs):
    ph(s, idx).insert_picture(str(img))
for idx, cap in zip((16, 17, 18), T['5']['captions']):
    set_text(ph(s, idx).text_frame, cap, size=13)
credit(s, T['5']['credit'])
notes(s, T['5']['notes'])

# 6 State of the art today — one big photo
s = prs.slides.add_slide(layout('Headline and 1 image'))
title(s, T['6']['title'])
ph(s, 13).insert_picture(str(IMG['destiny']))
credit(s, '© ESA – M. Wandt · ' + T['6']['caption'])
notes(s, T['6']['notes'])

# 7 What makes space stressful — same photo, prompt overlay
s = prs.slides.add_slide(layout('Empty'))
picture(s, IMG['destiny_small'], 0, 0, w=13.333)
rect(s, 0, 0, 13.333, 7.5, NAVY, transparency=55)
textbox(s, 1.0, 2.3, 11.3, 2.6, T['7']['prompt'], size=54, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
textbox(s, 1.0, 5.1, 11.3, 0.8, T['7']['sub'], size=20, color=RGBColor(0xDE, 0xF0, 0xFF), align=PP_ALIGN.CENTER)
credit(s, '© ESA – M. Wandt', x=9.9, y=7.05, w=3.3, color=RGBColor(0xDE, 0xF0, 0xFF))
notes(s, T['7']['notes'])

# 8 The cognitive link — diagram
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['8']['title'])
boxes = T['8']['boxes']  # [(head, sub), ...] three
bx = [0.9, 5.05, 9.2]; bw = 3.25; by = 2.5; bh = 2.3
fills = [SKY, KTH_BLUE, NAVY]
for i, (head, sub) in enumerate(boxes):
    b = rect(s, bx[i], by, bw, bh, fills[i], shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
    shape_text(b, [f"**{head}**", sub], size=16, color=NAVY if i == 0 else WHITE)
    if i < 2: arrow(s, bx[i] + bw + 0.08, by + bh / 2, bx[i + 1] - 0.08, by + bh / 2, color=NAVY, width=3)
textbox(s, 0.9, 5.25, 11.6, 1.2, T['8']['examples'], size=15, color=GREY)
notes(s, T['8']['notes'])

# 9 Why we care — navy statement
s = prs.slides.add_slide(layout('Title/chapter, navy blue'))
title(s, T['9']['title'])
set_text(ph(s, 1).text_frame, T['9']['sub'], size=20)
kill_ph(s, 13)
notes(s, T['9']['notes'])

# 10 Current stations — two images with captions
s = prs.slides.add_slide(layout('Headline and 2 images, text'))
title(s, T['10']['title'])
ph(s, 13).insert_picture(str(IMG['cupola'])); ph(s, 14).insert_picture(str(IMG['casa']))
set_text(ph(s, 16).text_frame, T['10']['captions'][0], size=14); set_text(ph(s, 17).text_frame, T['10']['captions'][1], size=14)
credit(s, '© ESA – M. Wandt')
notes(s, T['10']['notes'])

# 11 Future vision — Guidice artwork
s = prs.slides.add_slide(layout('Headline and 1 image'))
title(s, T['11']['title'])
ph(s, 13).insert_picture(str(IMG['guidice']))
credit(s, '© NASA – Rick Guidice, NASA Ames Research Center · ' + T['11']['caption'])
notes(s, T['11']['notes'])

# 12 Small steps today — two statement cards
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['12']['title'])
c1 = rect(s, 0.9, 2.3, 5.2, 3.4, LIGHT, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
shape_text(c1, [f"**{T['12']['left'][0]}**", T['12']['left'][1]], size=20, color=NAVY)
arrow(s, 6.25, 4.0, 7.15, 4.0, color=NAVY, width=4)
c2 = rect(s, 7.3, 2.3, 5.2, 3.4, NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
shape_text(c2, [f"**{T['12']['right'][0]}**", T['12']['right'][1]], size=20, color=WHITE)
textbox(s, 0.9, 6.0, 11.6, 0.6, T['12']['sub'], size=16, color=GREY, align=PP_ALIGN.CENTER)
notes(s, T['12']['notes'])

# 13 How we test cognition — three tools with captions
s = prs.slides.add_slide(layout('Headline and 3 images, text'))
title(s, T['13']['title'])
for idx, img in zip((13, 14, 15), (IMG['pvt'], IMG['nback'], SHOT / 'matb-run.png')):
    ph(s, idx).insert_picture(str(img))
for idx, cap in zip((16, 17, 18), T['13']['captions']):
    set_text(ph(s, idx).text_frame, cap, size=13)
textbox(s, 0.66, 6.55, 12, 0.4, T['13']['quote'], size=14, color=GREY)
notes(s, T['13']['notes'])

# 14 Where research runs — two images
s = prs.slides.add_slide(layout('Headline and 2 images, text'))
title(s, T['14']['title'])
ph(s, 13).insert_picture(str(IMG['mdrs_crew'])); ph(s, 14).insert_picture(str(IMG['wandt']))
set_text(ph(s, 16).text_frame, T['14']['captions'][0], size=14); set_text(ph(s, 17).text_frame, T['14']['captions'][1], size=14)
credit(s, '© ISAE-SUPAERO · © ESA – M. Wandt')
notes(s, T['14']['notes'])

# 15 Reaction Time — QR
qr_slide(T['15'], 'reaction-default', 'reaction-start', T['15']['bullets'])

# 16 MATB teacher demo — 2x2 grid of training screenshots
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['16']['title'])
shots16 = ('training-monitoring', 'training-tracking', 'training-comms', 'training-resource')
for i, (shot, cap) in enumerate(zip(shots16, T['16']['captions'])):
    gx = 0.66 + (i % 2) * 6.2; gy = 2.0 + (i // 2) * 2.45
    pic = fit_picture(s, SHOT / f'{shot}.png', gx, gy, 3.3, 2.0)
    pic.line.color.rgb = RGBColor(0xC8, 0xD0, 0xDA); pic.line.width = Pt(0.75)
    textbox(s, gx + 3.4, gy + 0.3, 2.7, 1.6, cap, size=13, color=NAVY)
textbox(s, 0.66, 6.85, 12, 0.4, T['16']['sub'], size=13, color=GREY)
notes(s, T['16']['notes'])

# 17 MATB dry run — QR
qr_slide(T['17'], 'matb-2min', 'matb-end', T['17']['bullets'])

# 17b Condition Lab — QR
qr_slide(T['17b'], 'condition-lab', 'condition-lab-result', T['17b']['bullets'])

# 18 Learning & plateau — native line chart
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['18']['title'])
cd = CategoryChartData()
cd.categories = T['18']['x']
cd.add_series(T['18']['series'], (52, 61, 68, 73, 76, 75, 78, 76, 77, 76))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.66), Inches(2.0), Inches(7.6), Inches(4.8), cd)
ch = gf.chart; ch.has_legend = False; ch.has_title = False
ch.plots[0].series[0].format.line.color.rgb = KTH_BLUE; ch.plots[0].series[0].format.line.width = Pt(2.5)
ch.plots[0].series[0].smooth = True
ch.value_axis.has_major_gridlines = True; ch.value_axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xDE, 0xE5)
ch.value_axis.minimum_scale = 40; ch.value_axis.maximum_scale = 90
ch.value_axis.tick_labels.font.size = Pt(11); ch.category_axis.tick_labels.font.size = Pt(11)
ch.value_axis.tick_labels.font.color.rgb = GREY; ch.category_axis.tick_labels.font.color.rgb = GREY
ch.category_axis.format.line.color.rgb = RGBColor(0xB0, 0xB8, 0xC4); ch.value_axis.format.line.fill.background()
ch.value_axis.has_title = True; ch.value_axis.axis_title.text_frame.text = T['18']['y']
ch.value_axis.axis_title.text_frame.paragraphs[0].runs[0].font.size = Pt(11)
textbox(s, 8.6, 2.1, 4.1, 4.6, [(b, 0) for b in T['18']['bullets']], size=16, color=NAVY, space_after=10)
notes(s, T['18']['notes'])

# 19 Simulator — QR
qr_slide(T['19'], 'simulator', 'simulator', T['19']['bullets'])
# 20 Blueprint — QR
qr_slide(T['20'], 'blueprint', 'blueprint-scenario', T['20']['bullets'], bullet_size=16)
# 21 Model Lab — QR
qr_slide(T['21'], 'model-lab', 'model-lab-map', T['21']['bullets'])

# 22 Exit ticket — two sticky notes
s = prs.slides.add_slide(layout('Headline only'))
title(s, T['22']['title'])
textbox(s, 0.66, 1.95, 12, 0.5, T['22']['sub'], size=16, color=GREY)
for i, (lab, body) in enumerate(T['22']['notes_cards']):
    x = 1.6 + i * 5.6
    n_ = rect(s, x, 2.7, 4.5, 3.6, YELLOW if i == 0 else RGBColor(0xFF, 0xC9, 0xB0))
    n_.rotation = -2 if i == 0 else 2
    n_.shadow.inherit = True
    shape_text(n_, [f"**{lab}**", body], size=20, color=NAVY)
notes(s, T['22']['notes'])

# 23 Closing messages — navy
s = prs.slides.add_slide(layout('Title/chapter, navy blue'))
kill_ph(s, 0); kill_ph(s, 1); kill_ph(s, 13)
textbox(s, 0.66, 1.9, 11.5, 1.5, T['23']['title'], size=40, color=WHITE, bold=True, anchor=MSO_ANCHOR.BOTTOM)
textbox(s, 0.66, 3.7, 11.5, 3.2, [(f"**{m}**", 0) for m in T['23']['messages']], size=22, color=WHITE, space_after=10)
notes(s, T['23']['notes'])

# 24 Thanks — navy with hub QR + logos
s = prs.slides.add_slide(layout('Title/chapter, navy blue'))
kill_ph(s, 13)
set_text(ph(s, 0).text_frame, T['24']['title'], size=40)
set_text(ph(s, 1).text_frame, [T['24']['url'], T['24']['sub']], size=18)
r = rect(s, 9.6, 1.9, 3.0, 3.0, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
picture(s, QR / f'hub-{LANG}.png', 9.8, 2.1, 2.6, 2.6)
textbox(s, 9.6, 4.95, 3.0, 0.4, T['24']['scan'], size=12, color=WHITE, align=PP_ALIGN.CENTER)
picture(s, IMG['esero'], 9.6, 5.75, h=0.65)
r = rect(s, 11.7, 5.67, 0.9, 0.8, WHITE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
picture(s, IMG['kth'], 11.78, 5.71, h=0.72)
notes(s, T['24']['notes'])

# Greek: the KTH heading font has no Greek glyphs (falls back to a serif) — use Arial throughout
if LANG == 'el':
    from pptx.util import Pt as _Pt
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                is_title = sh.is_placeholder and sh.placeholder_format.idx == 0
                lay = sl.slide_layout.name
                for para in sh.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name = 'Arial'
                        if is_title and lay == 'Text and image':
                            run.font.size = _Pt(28)
                        elif is_title and lay.startswith('Title/chapter'):
                            run.font.size = _Pt(36)
            if getattr(sh, 'has_chart', False) and sh.has_chart:
                pass

# prune layouts this deck does not use (keeps the file small and the layout picker tidy)
used = {sl.slide_layout.name for sl in prs.slides} | {'Title, KTH blue', 'Headline and content', 'Two parts', 'Comparison, KTH blue'}
for lay in list(prs.slide_master.slide_layouts):
    if lay.name not in used:
        try:
            prs.slide_master.slide_layouts.remove(lay)
        except Exception as e:
            print('could not remove layout', lay.name, e)
prs.save(OUT)
print('saved', OUT, len(prs.slides), 'slides')
